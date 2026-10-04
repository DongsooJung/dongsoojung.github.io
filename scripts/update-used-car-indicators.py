#!/usr/bin/env python3
"""Publish monthly vehicle transfer registrations from official MOLIT XLSX files.

These registrations are an indirect measure of used-car activity. They are not
advertised listings, completed used-car sales, or transaction prices.
"""

from __future__ import annotations

import csv
import io
import json
import re
from datetime import datetime, timezone
from pathlib import Path

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook


SOURCE_URL = "https://stat.molit.go.kr/portal/cate/statMetaView.do?hRsId=58"
DOWNLOAD_URL = "https://stat.molit.go.kr/portal/common/downLoadFile.do"
OUTPUT_DIR = Path("strategy/used-car/data")
JSON_PATH = OUTPUT_DIR / "market-indicators.json"
CSV_PATH = OUTPUT_DIR / "market-indicators.csv"
MONTH_COUNT = 12
FILE_PATTERN = re.compile(r"(?P<year>\d{4})년\s*(?P<month>\d{1,2})월\s*자동차\s*등록자료\s*통계\.xlsx")
EXPECTED_HEADERS = ("계", "업자매매", "당사자매매", "증여", "상속", "촉탁", "기타")
REGION_NAMES = {
    "서울", "부산", "대구", "인천", "대전", "울산", "세종",
    "경기", "충북", "충남", "경북", "경남", "제주", "강원", "전북",
}
SPLIT_JEONNAM_GWANGJU = REGION_NAMES | {"전남", "광주"}
COMBINED_JEONNAM_GWANGJU = REGION_NAMES | {"전남광주"}


def to_int(value: object, context: str) -> int:
    if isinstance(value, bool):
        raise ValueError(f"{context}: invalid boolean value")
    if isinstance(value, int):
        result = value
    elif isinstance(value, float) and value.is_integer():
        result = int(value)
    elif isinstance(value, str) and re.fullmatch(r"[\d,]+", value.strip()):
        result = int(value.replace(",", ""))
    else:
        raise ValueError(f"{context}: invalid count {value!r}")
    if result < 0:
        raise ValueError(f"{context}: negative count")
    return result


def discover_files(session: requests.Session) -> list[tuple[str, str]]:
    response = session.get(SOURCE_URL, timeout=30)
    response.raise_for_status()
    soup = BeautifulSoup(response.text, "html.parser")
    files: dict[str, str] = {}
    for anchor in soup.select('a[href="/portal/common/downLoadFile.do"]'):
        match = re.search(r"downFile\('([^']+)'", anchor.get("onclick", ""))
        if not match:
            continue
        filename = match.group(1)
        period_match = FILE_PATTERN.fullmatch(filename)
        if not period_match:
            continue
        year, month = int(period_match["year"]), int(period_match["month"])
        if not 1 <= month <= 12:
            continue
        period = f"{year:04d}-{month:02d}"
        files.setdefault(period, filename)
    latest = sorted(files.items(), reverse=True)[:MONTH_COUNT]
    if len(latest) != MONTH_COUNT:
        raise ValueError(f"Expected {MONTH_COUNT} monthly files, found {len(latest)}")
    for (newer, _), (older, _) in zip(latest, latest[1:]):
        newer_num = int(newer[:4]) * 12 + int(newer[5:])
        older_num = int(older[:4]) * 12 + int(older[5:])
        if newer_num - older_num != 1:
            raise ValueError(f"Missing month between {older} and {newer}")
    return list(reversed(latest))


def parse_file(session: requests.Session, period: str, filename: str) -> dict:
    response = session.get(
        DOWNLOAD_URL,
        params={"oFileName": filename, "rFileName": filename, "midpath": "/stat_file/"},
        timeout=(10, 60),
    )
    response.raise_for_status()
    if not response.content.startswith(b"PK\x03\x04") or len(response.content) > 15_000_000:
        raise ValueError(f"{period}: expected a valid XLSX file")
    workbook = load_workbook(io.BytesIO(response.content), read_only=True, data_only=True)
    try:
        candidates = [sheet for sheet in workbook.worksheets if "이전" in sheet.title and "당월" in sheet.title]
        if len(candidates) != 1:
            raise ValueError(f"{period}: expected one monthly transfer worksheet, found {len(candidates)}")
        sheet = candidates[0]
        headers = tuple(str(sheet.cell(3, col).value or "").strip() for col in range(3, 10))
        if headers != EXPECTED_HEADERS:
            raise ValueError(f"{period}: unexpected worksheet headers: {headers}")

        regions = []
        region_counts = []
        total = None
        for row in sheet.iter_rows(min_row=5, max_col=9, values_only=True):
            name = str(row[0] or "").strip()
            if name == "총계":
                total = [to_int(row[col], f"{period} total col {col + 1}") for col in range(2, 9)]
                break
            if not name:
                continue
            if name not in SPLIT_JEONNAM_GWANGJU | COMBINED_JEONNAM_GWANGJU:
                raise ValueError(f"{period}: unexpected region {name!r}")
            counts = [to_int(row[col], f"{period} {name} col {col + 1}") for col in range(2, 9)]
            if counts[0] != sum(counts[1:]):
                raise ValueError(f"{period} {name}: components do not add to total")
            region_counts.append(counts)
            regions.append({
                "name": name,
                "totalTransfers": counts[0],
                "dealerTransfers": counts[1],
                "privateTransfers": counts[2],
            })
        names = {region["name"] for region in regions}
        if total is None or names not in (SPLIT_JEONNAM_GWANGJU, COMBINED_JEONNAM_GWANGJU):
            raise ValueError(f"{period}: missing total or region rows")
        if total[0] != sum(total[1:]):
            raise ValueError(f"{period}: total components do not add up")
        if total != [sum(row[col] for row in region_counts) for col in range(7)]:
            raise ValueError(f"{period}: region sums do not match national total")
        return {
            "period": period,
            "totalTransfers": total[0],
            "dealerTransfers": total[1],
            "privateTransfers": total[2],
            "otherTransfers": sum(total[3:]),
            "regions": sorted(regions, key=lambda item: item["totalTransfers"], reverse=True),
        }
    finally:
        workbook.close()


def write_snapshot(months: list[dict]) -> None:
    latest = months[-1]
    prior = months[-2]
    payload = {
        "source": "국토교통부 자동차등록현황보고 — 이전 등록현황(당월)",
        "sourceUrl": SOURCE_URL,
        "updatedAt": datetime.now(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z"),
        "period": latest["period"],
        "totalTransfers": latest["totalTransfers"],
        "dealerTransfers": latest["dealerTransfers"],
        "privateTransfers": latest["privateTransfers"],
        "otherTransfers": latest["otherTransfers"],
        "dealerShare": round(latest["dealerTransfers"] / latest["totalTransfers"] * 100, 1),
        "monthOverMonth": round((latest["totalTransfers"] / prior["totalTransfers"] - 1) * 100, 1),
        "trend": [{key: month[key] for key in ("period", "totalTransfers", "dealerTransfers", "privateTransfers", "otherTransfers")} for month in months],
        "regions": latest["regions"],
        "methodology": "국토교통부 월별 자동차 이전등록 건수. 실제 판매 완료 건수, 현재 매물 수, 거래가격을 뜻하지 않는 간접 지표입니다. 업자매매·당사자매매 외에는 증여·상속·촉탁·기타가 포함됩니다.",
        "regionNote": (
            "해당 월 원본 통계에서 전남과 광주는 '전남광주'로 합산되어 있습니다."
            if any(region["name"] == "전남광주" for region in latest["regions"])
            else "해당 월 원본 통계에서 전남과 광주는 별도 집계되어 있습니다."
        ),
    }
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    csv_buffer = io.StringIO(newline="")
    columns = ("period", "totalTransfers", "dealerTransfers", "privateTransfers", "otherTransfers")
    writer = csv.DictWriter(csv_buffer, fieldnames=columns, lineterminator="\n")
    writer.writeheader()
    for month in months:
        writer.writerow({key: month[key] for key in columns})
    # Build both outputs before replacing either existing snapshot. A source error
    # above this point leaves the previously published files untouched.
    csv_text = csv_buffer.getvalue()
    if JSON_PATH.exists() and CSV_PATH.exists():
        previous = json.loads(JSON_PATH.read_text(encoding="utf-8"))
        previous.pop("updatedAt", None)
        current = dict(payload)
        current.pop("updatedAt", None)
        if previous == current and CSV_PATH.read_text(encoding="utf-8-sig") == csv_text:
            print("Official data unchanged; keeping the existing snapshot")
            return
    JSON_PATH.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    CSV_PATH.write_text(csv_text, encoding="utf-8-sig")


def main() -> None:
    with requests.Session() as session:
        session.headers.update({"User-Agent": "STARGATE-UsedCarIndicators/1.0 (public statistics)"})
        files = discover_files(session)
        months = []
        for period, filename in files:
            month = parse_file(session, period, filename)
            months.append(month)
            print(f"{period}: {month['totalTransfers']:,} transfers")
    write_snapshot(months)
    print(f"Published {len(months)} months through {months[-1]['period']}")


if __name__ == "__main__":
    main()
